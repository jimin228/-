# -*- coding: utf-8 -*-
"""
외근대장 Excel 자동 생성 스크립트 (1팀 ~ 4팀)
- openpyxl을 사용하여 깔끔한 결재란과 팀별 시트가 포함된 엑셀 파일을 생성합니다.
"""

import os
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl 모듈이 필요합니다. 'pip install openpyxl'을 실행해주세요.")
    exit(1)


def create_outwork_workbook(filename="외근대장_4개팀_통합.xlsx"):
    wb = openpyxl.Workbook()
    # 기본 첫 번째 시트
    default_sheet = wb.active
    wb.remove(default_sheet)

    # 4개 팀 정보
    teams = [
        {"id": "all", "title": "전체 통합 외근대장", "sheet_name": "전체_통합"},
        {"id": "1", "title": "1팀 외근대장", "sheet_name": "1팀"},
        {"id": "2", "title": "2팀 외근대장", "sheet_name": "2팀"},
        {"id": "3", "title": "3팀 외근대장", "sheet_name": "3팀"},
        {"id": "4", "title": "4팀 외근대장", "sheet_name": "4팀"},
    ]

    # 스타일 정의
    font_title = Font(name="맑은 고딕", size=18, bold=True, color="1F2937")
    font_header = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="맑은 고딕", size=10, color="111827")
    font_small = Font(name="맑은 고딕", size=9, color="6B7280")
    font_approval = Font(name="맑은 고딕", size=9, bold=True, color="374151")

    fill_header = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    fill_approval_header = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    headers = [
        ("No.", 6, align_center),
        ("외근일자", 12, align_center),
        ("소속팀", 10, align_center),
        ("성명", 10, align_center),
        ("직급", 10, align_center),
        ("방문처 (행선지)", 25, align_left),
        ("외근 목적", 30, align_left),
        ("출발시간", 11, align_center),
        ("복귀시간", 11, align_center),
        ("교통편", 12, align_center),
        ("동행자", 14, align_center),
        ("진행상태", 10, align_center),
        ("비고 / 연락처", 22, align_left),
    ]

    for team_info in teams:
        ws = wb.create_sheet(title=team_info["sheet_name"])
        ws.views.sheetView[0].showGridLines = True

        # 1. 대장 제목 (Row 2, Column B)
        ws.merge_cells("B2:G2")
        title_cell = ws["B2"]
        title_cell.value = team_info["title"]
        title_cell.font = font_title
        title_cell.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[2].height = 36

        # 작성일자
        ws["B3"] = f"작성일자: {datetime.now().strftime('%Y-%m-%d')} 기준"
        ws["B3"].font = font_small

        # 2. 결재란 (Row 2~3, Column K~N)
        # K2:K3 병합 (결재 텍스트)
        ws.merge_cells("K2:K3")
        k2 = ws["K2"]
        k2.value = "결\n재"
        k2.font = font_approval
        k2.fill = fill_approval_header
        k2.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        appr_cols = [("L", "담당"), ("M", "팀장"), ("N", "부서장")]
        for col_letter, label in appr_cols:
            c_header = ws[f"{col_letter}2"]
            c_header.value = label
            c_header.font = font_approval
            c_header.fill = fill_approval_header
            c_header.alignment = align_center

            # 서명란 높이
            ws[f"{col_letter}3"].value = ""

        ws.row_dimensions[3].height = 32

        # 결재란 테두리 적용
        for r in range(2, 4):
            for c in range(11, 15):
                cell = ws.cell(row=r, column=c)
                cell.border = thin_border

        # 3. 테이블 헤더 (Row 5)
        ws.row_dimensions[5].height = 28
        start_col = 2  # B열부터 시작
        for idx, (h_name, width, _) in enumerate(headers):
            col_idx = start_col + idx
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=5, column=col_idx)
            cell.value = h_name
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border
            ws.column_dimensions[col_letter].width = width

        # A열은 여백
        ws.column_dimensions["A"].width = 3

        # 4. 빈 데이터 행 서식 생성 (Row 6 ~ 25)
        for r in range(6, 26):
            ws.row_dimensions[r].height = 22
            row_no = r - 5
            for idx, (_, _, default_align) in enumerate(headers):
                col_idx = start_col + idx
                cell = ws.cell(row=r, column=col_idx)
                cell.border = thin_border
                cell.font = font_body
                cell.alignment = default_align

                if idx == 0:  # No. 번호
                    cell.value = row_no
                    cell.alignment = align_center
                elif idx == 2 and team_info["id"] != "all":
                    cell.value = f"{team_info['id']}팀"

                # 교차 행 배경색 (얼룩말 무늬)
                if r % 2 == 1:
                    cell.fill = fill_zebra

    # 저장
    wb.save(filename)
    print(f"✅ 외근대장 Excel 파일이 성공적으로 생성되었습니다: {filename}")


if __name__ == "__main__":
    current_dir = os.path.dirname(os.path.abspath(__file__))
    target_path = os.path.join(current_dir, "외근대장_4개팀_통합.xlsx")
    create_outwork_workbook(target_path)
